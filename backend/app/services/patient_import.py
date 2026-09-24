"""
模块名称：patient_import.py
所属层级：服务层（services）
功能说明：解析患者批量导入文件，统一输出逐行原始字段；业务校验与落库由接口层完成。

安全边界：
    - 只读取内存中的 csv/xlsx 内容，不执行公式、不访问外部链接；
    - xlsx 以 read_only + data_only 打开；
    - 文件大小与扩展名由接口层先行限制。

修改历史：
    - 2026-09-24  v1.0  第二轮第二批新增
"""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_HEADERS: tuple[str, ...] = ("姓名", "出生年", "出生月", "性别", "住院号", "当前科室", "联系方式")


def _normalize_rows(rows: list[list[object]]) -> list[tuple[int, dict[str, object]]]:
    """把二维表按中文表头转换为“源文件行号 + 字段字典”。"""
    if not rows:
        raise ValueError("导入文件没有内容。")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError(f"导入文件缺少列：{'、'.join(missing)}。")
    indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}
    parsed: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        # 全空行属于排版空白，不计入成功、跳过或失败
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        parsed.append(
            (
                row_number,
                {
                    header: values[indexes[header]] if indexes[header] < len(values) else None
                    for header in REQUIRED_HEADERS
                },
            )
        )
    return parsed


def parse_patient_import(filename: str, content: bytes) -> list[tuple[int, dict[str, object]]]:
    """按扩展名解析 UTF-8 csv 或首个工作表 xlsx。"""
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 文件必须使用 UTF-8 编码。") from exc
        return _normalize_rows([list(row) for row in csv.reader(StringIO(text))])
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("无法读取 xlsx 文件，请确认文件未损坏。") from exc
        try:
            worksheet = workbook.worksheets[0]
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return _normalize_rows(rows)
    raise ValueError("仅支持 .xlsx 或 .csv 文件。")
